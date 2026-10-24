"""Genotype file parsers. Every parser returns the same `Genotypes` shape:
calls[i][j] is the genotype of sample i at SNP j as a sorted 2-letter string ("AG") or None."""
import csv
import io
from dataclasses import dataclass
from pathlib import Path

import openpyxl

MISSING = {"", "?", "N", "NN", "00", "--", "0", "NONE", "NA"}


@dataclass
class Genotypes:
    samples: list[str]
    snps: list[str]
    calls: list[list[str | None]]

    def call_rate(self, i: int) -> float:
        row = self.calls[i]
        return sum(c is not None for c in row) / max(len(row), 1)


def norm_call(v) -> str | None:
    """'A:G', 'A G', 'GA', 'a/g' -> 'AG'; anything unreadable -> None."""
    if v is None:
        return None
    s = "".join(ch for ch in str(v).upper() if ch.isalnum())
    if s in MISSING or len(s) != 2 or not set(s) <= set("ACGT"):
        return None
    return "".join(sorted(s))


def parse_hapmap(path: str | Path) -> Genotypes:
    rows = [line.rstrip("\r\n").split("\t") for line in open(path)]
    samples, snps, cols = rows[0][11:], [], []
    for r in rows[1:]:
        snps.append(r[0])
        cols.append([norm_call(x) for x in r[11:]])
    calls = [list(c) for c in zip(*cols)]  # SNP-major -> sample-major
    return Genotypes(samples, snps, calls)


def read_table(raw: bytes) -> list[list[str]]:
    """Rows of a CSV saved with any delimiter (Excel in an Indonesian locale writes ';') or of an .xlsx's
    first sheet. Raises UnicodeDecodeError for anything else."""
    if raw[:2] == b"PK":  # .xlsx is a zip archive
        ws = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        return [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    text = raw.decode("utf-8-sig")
    head = text.split("\n", 1)[0]
    return list(csv.reader(io.StringIO(text), delimiter=max(",;\t", key=head.count)))


def parse_kasp_csv(text: str) -> Genotypes:
    return parse_kasp_rows(read_table(text.encode()))


def parse_kasp_rows(rows: list[list[str]]) -> Genotypes:
    """Lab export: first column sample id, one column per SNP."""
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        raise ValueError("File kosong atau tidak punya baris data")
    snps = [s.strip() for s in rows[0][1:]]
    samples = [r[0].strip() for r in rows[1:]]
    calls = [[norm_call(v) for v in (r[1:] + [""] * len(snps))[: len(snps)]] for r in rows[1:]]
    return Genotypes(samples, snps, calls)


def parse_colombia_table4(path: str | Path) -> tuple[Genotypes, list[str]]:
    """Osorio-Guarin 2025 Supplementary Table 4. Returns genotypes and the genebank of each tree."""
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    snps = [s for s in rows[1][2:] if s]
    samples, banks, calls = [], [], []
    for r in rows[2:]:
        if not r[0]:
            continue
        samples.append(str(r[0]).strip())
        banks.append(str(r[1]).strip())
        calls.append([norm_call(v) for v in r[2:2 + len(snps)]])
    return Genotypes(samples, snps, calls), banks
